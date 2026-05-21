import tkinter as tk    
from tkinter import ttk
import openpyxl as azl
from tkinter import messagebox


wb = azl.Workbook()
ws = wb.active
ws.append(["Order ID", "Customer Name", "Product", "Quantity", "Price", "Total"])    
record_id = 1


def endozo():
    global record_id
    cname = cname_entry.get()
    product = product_entry.get()
    qty = qty_entry.get()
    price = price_entry.get()


    if not (cname and product and qty and price):
        tk.messagebox.showerror("Input Error", "Please fill in all fields.")
        return


    try:
        qty = int(qty)
        price = float(price)
    except ValueError:
        tk.messagebox.showerror("Input Error", "Quantity must be an integer and Price must be a number.")
        return


    total = qty * price
    table.insert("", "end", values=(record_id, cname, product, qty, price, total))
    ws.append([record_id, cname, product, qty, price, total])
    wb.save("ordersDB.xlsx")
    record_id += 1
    clear_entries()


def clear_entries():
    cname_entry.delete(0, tk.END)
    product_entry.delete(0, tk.END)
    qty_entry.delete(0, tk.END)
    price_entry.delete(0, tk.END)


def update_record():
    selected_item = table.selection()
    if not selected_item:
        tk.messagebox.showerror("Selection Error", "Please select a record to update.")
        return


    item = table.item(selected_item)
    record_id = item['values'][0]
    cname = cname_entry.get()
    product = product_entry.get()
    qty = qty_entry.get()
    price = price_entry.get()
    if not (cname and product and qty and price):
        tk.messagebox.showerror("Input Error", "Please fill in all fields.")
        return
   
    try:
        qty = int(qty)
        price = float(price)
    except ValueError:
        tk.messagebox.showerror("Input Error", "Quantity must be an integer and Price must be a number.")
        return
    total = qty * price
    table.item(selected_item, values=(record_id, cname, product, qty, price,
total))
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == record_id:
            ws.delete_rows(row[0] + 1)
            ws.append([record_id, cname, product, qty, price, total])
            break      
    wb.save("ordersDB.xlsx")


def delete_record():
    selected_item = table.selection()
    if not selected_item:
        tk.messagebox.showerror("Selection Error", "Please select a record to delete.")
        return


    item = table.item(selected_item)
    record_id = item['values'][0]
    table.delete(selected_item)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == record_id:
            ws.delete_rows(row[0] + 1)
            break
    wb.save("ordersDB.xlsx")


def autopopulate_entries(event):
    selected_item = table.selection()
    if selected_item:
        item = table.item(selected_item)
        values = item['values']
        cname_entry.delete(0, tk.END)
        cname_entry.insert(0, values[1])
        product_entry.delete(0, tk.END)
        product_entry.insert(0, values[2])
        qty_entry.delete(0, tk.END)
        qty_entry.insert(0, values[3])
        price_entry.delete(0, tk.END)
        price_entry.insert(0, values[4])


window = tk.Tk()
window.title("Simple Ordering System")
window.configure(bg="lightblue")


# Form Title
title = tk.Label(window, text="Simple Ordering System", font=("Times New Roman", 14, "bold"), bg="lightblue")
title.grid(row=0, column=0, columnspan=6)


# Frame
genframe = tk.Frame(window, bg="lightblue", bd=2, relief="groove")
genframe.grid(row=1, column=0, columnspan=7, padx=10, pady=10)


# Customer Name Entry
cname_entry = tk.Entry(genframe, font=("Poppins", 12))
cname_entry.grid(row=2, column=1, columnspan=2, padx=10, pady=(10, 0))


cname_label = tk.Label(genframe, text="Customer Name", font=("Poppins", 10, "italic"), bg="lightblue")
cname_label.grid(row=3, column=1, columnspan=2)


# Product Entry
product_entry = tk.Entry(genframe, font=("Poppins", 12))
product_entry.grid(row=2, column=3, columnspan=2, padx=10, pady=(10, 0))


product_label = tk.Label(genframe, text="Product", font=("Poppins", 10, "italic"), bg="lightblue")
product_label.grid(row=3, column=3, columnspan=2)


# Quantity Entry
qty_entry = tk.Entry(genframe, font=("Poppins", 12))
qty_entry.grid(row=4, column=1, columnspan=2, padx=10, pady=(10, 0))


qty_label = tk.Label(genframe, text="Quantity", font=("Poppins", 10, "italic"), bg="lightblue")
qty_label.grid(row=5, column=1, columnspan=2)


# Price Entry
price_entry = tk.Entry(genframe, font=("Poppins", 12))
price_entry.grid(row=4, column=3, columnspan=2, padx=10, pady=(10, 0))


price_label = tk.Label(genframe, text="Price", font=("Poppins", 10, "italic"), bg="lightblue")
price_label.grid(row=5, column=3, columnspan=2)


# Buttons
submit_btn = tk.Button(window, text="Submit", font=("Poppins", 12, "bold"), bg="lightpink", command=endozo)
submit_btn.grid(row=6, column=1, pady=(10, 20))


update_btn = tk.Button(window, text="Update",font=("Poppins", 12, "bold"), bg="lightgreen")
update_btn.grid(row=6, column=2)


delete_btn = tk.Button(window, text="Delete", bg="red", fg="white",font=("Poppins", 12, "bold"))
delete_btn.grid(row=6, column=3)


# Table
table = ttk.Treeview(
    window,
    columns=("Order ID", "Customer Name", "Product", "Quantity", "Price", "Total"),
    show="headings"
)


for headings in ("Order ID", "Customer Name", "Product", "Quantity", "Price", "Total"):
    table.heading(headings, text=headings)


table.grid(row=7, column=0, columnspan=6, padx=10, pady=10)
 
def load_data():
    for row in ws.iter_rows(min_row=2, values_only=True):
        table.insert("", "end", values=row)
load_data()
update_btn.config(command=update_record)
delete_btn.config(command=delete_record)


def on_row_select(event):
    autopopulate_entries(event)
table.bind("<<TreeviewSelect>>", on_row_select)






























window.mainloop()



