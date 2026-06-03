import tkinter as azl
from tkinter import ttk
from tkinter import messagebox
import openpyxl as op
import os

# Create database if it doesn't exist
if not os.path.exists("Endozo_Database.xlsx"):
    wb = op.Workbook()
    ws = wb.active
    ws.append(["Order ID", "Customer Name", "Month", "Day", "Time"])
    wb.save("Endozo_Database.xlsx")

window = azl.Tk()
window.title("Gym Membership System")
window.configure(bg="#dea3bd")


def refresh_table():
    for row in table.get_children():
        table.delete(row)

    workbook = op.load_workbook("Endozo_Database.xlsx")
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2, values_only=True):
        table.insert("", "end", values=row)


def rss():
    if not mname_entry.get() or not month_entry.get() or not day_entry.get() or not time_entry.get():
        messagebox.showwarning("Input Error", "All fields are required")
        return

    # Month Validation
    if not month_entry.get().isdigit():
        messagebox.showerror("Error", "Month must be a number")
        return

    month = int(month_entry.get())

    if month < 1 or month > 12:
        messagebox.showerror("Error", "Invalid month")
        return

    # Day Validation
    if not day_entry.get().isdigit():
        messagebox.showerror("Error", "Day must be a number")
        return

    day = int(day_entry.get())

    if day < 1 or day > 31:
        messagebox.showerror("Error", "Invalid day")
        return

    # Time Validation
    try:
        hour, minute = map(int, time_entry.get().split(":"))

        if hour < 0 or hour > 23 or minute < 0 or minute > 59:
            raise ValueError

    except:
        messagebox.showerror(
            "Error",
            "Use 24-hour format like 13:30"
        )
        return

    workbook = op.load_workbook("Endozo_Database.xlsx")
    sheet = workbook.active

    # Duplicate Check
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if (
            str(row[2]) == month_entry.get() and
            str(row[3]) == day_entry.get() and
            str(row[4]) == time_entry.get()
        ):
            messagebox.showerror(
                "Duplicate Member",
            "Member already exists."
            )
            return

    # Generate Unique ID
    new_id = 1

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0]:
            new_id = max(new_id, int(row[0]) + 1)

    sheet.append([
        new_id,
        mname_entry.get(),
        month_entry.get(),
        day_entry.get(),
        time_entry.get()
    ])

    workbook.save("Endozo_Database.xlsx")

    refresh_table()

    mname_entry.delete(0, azl.END)
    month_entry.delete(0, azl.END)
    day_entry.delete(0, azl.END)
    time_entry.delete(0, azl.END)

    messagebox.showinfo("Success", "Membership saved successfully")


def rssl(event):
    selected = table.focus()

    if not selected:
        return

    values = table.item(selected, "values")

    if not values:
        return

    mname_entry.delete(0, azl.END)
    month_entry.delete(0, azl.END)
    day_entry.delete(0, azl.END)
    time_entry.delete(0, azl.END)

    mname_entry.insert(0, values[1])
    month_entry.insert(0, values[2])
    day_entry.insert(0, values[3])
    time_entry.insert(0, values[4])


def ndz():
    selected_item = table.selection()

    if not selected_item:
        messagebox.showwarning("No Selection", "Select a record first")
        return

    item = table.item(selected_item)

    order_id = item["values"][0]

    workbook = op.load_workbook("Endozo_Database.xlsx")
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2):
        if str(row[0].value) == str(order_id):
            row[1].value = mname_entry.get()
            row[2].value = month_entry.get()
            row[3].value = day_entry.get()
            row[4].value = time_entry.get()
            break

    workbook.save("Endozo_Database.xlsx")

    refresh_table()

    messagebox.showinfo("Updated", "Membership updated successfully")

    mname_entry.delete(0, azl.END)
    month_entry.delete(0, azl.END)
    day_entry.delete(0, azl.END)
    time_entry.delete(0, azl.END)


def donny():
    selected_item = table.selection()

    if not selected_item:
        messagebox.showwarning("No Selection", "Select a record first")
        return

    item = table.item(selected_item)

    order_id = item["values"][0]

    workbook = op.load_workbook("Endozo_Database.xlsx")
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2):
        if str(row[0].value) == str(order_id):
            sheet.delete_rows(row[0].row, 1)
            break

    workbook.save("Endozo_Database.xlsx")

    refresh_table()

    messagebox.showinfo("Deleted", "Membership deleted successfully")

    mname_entry.delete(0, azl.END)
    month_entry.delete(0, azl.END)
    day_entry.delete(0, azl.END)
    time_entry.delete(0, azl.END)


# Title
title = azl.Label(
    window,
    text="Gym Membership System",
    font=("Times New Roman", 14, "bold"),
    bg="#dea3bd"
)
title.grid(row=0, column=0, columnspan=6)

# Frame
genframe = azl.Frame(window, bg="#d96f9e", bd=2, relief="groove")
genframe.grid(row=1, column=0, columnspan=7, padx=10, pady=10)

# Members Name
mname_entry = azl.Entry(genframe, font=("Poppins", 12))
mname_entry.grid(row=2, column=1, columnspan=2, padx=10, pady=(10, 0))

mname_label = azl.Label(
    genframe,
    text="Members Name",
    font=("Poppins", 10, "italic"),
    bg="#d96f9e"
)
mname_label.grid(row=3, column=1, columnspan=2)

# Month
month_entry = azl.Entry(genframe, font=("Poppins", 12))
month_entry.grid(row=2, column=3, columnspan=2, padx=10, pady=(10, 0))

month_label = azl.Label(
    genframe,
    text="Month",
    font=("Poppins", 10, "italic"),
    bg="#d96f9e"
)
month_label.grid(row=3, column=3, columnspan=2)

# Day
day_entry = azl.Entry(genframe, font=("Poppins", 12))
day_entry.grid(row=4, column=1, columnspan=2, padx=10, pady=(10, 0))

day_label = azl.Label(
    genframe,
    text="Day",
    font=("Poppins", 10, "italic"),
    bg="#d96f9e"
)
day_label.grid(row=5, column=1, columnspan=2)

# Time
time_entry = azl.Entry(genframe, font=("Poppins", 12))
time_entry.grid(row=4, column=3, columnspan=2, padx=10, pady=(10, 0))

time_label = azl.Label(
    genframe,
    text="Time",
    font=("Poppins", 10, "italic"),
    bg="#d96f9e"
)
time_label.grid(row=5, column=3, columnspan=2)

# Buttons
submit_btn = azl.Button(
    window,
    text="Submit",
    width=10,
    font=("Times New Roman", 12, "bold"),
    bg="#769174",
    fg="white",
    command=rss
)
submit_btn.grid(row=6, column=1, padx=10, pady=10)

update_btn = azl.Button(
    window,
    text="Update",
    width=10,
    font=("Times New Roman", 12, "bold"),
    bg="#516d9f",
    fg="white",
    command=rssl
)
update_btn.grid(row=6, column=2, padx=10, pady=10)

delete_btn = azl.Button(
    window,
    text="Delete",
    width=10,
    font=("Times New Roman", 12, "bold"),
    bg="#990000",
    fg="white",
    command=donny
)
delete_btn.grid(row=6, column=3, padx=10, pady=10)

# Table
table = ttk.Treeview(
    window,
    columns=("Order ID", "Customer Name", "Month", "Day", "Time"),
    show="headings"
)

for heading in ("Order ID", "Customer Name", "Month", "Day", "Time"):
    table.heading(heading, text=heading)

table.grid(row=7, column=0, columnspan=6, padx=10, pady=10)
table.bind("<ButtonRelease-1>", rss)

# Load existing data
refresh_table()

window.mainloop()